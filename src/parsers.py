"""
parsers.py - STAGE 1 of the tool: turn the input files into clean, structured text.

    parse_rfp(path)  -> the RFP (PDF or Word) as a list of "blocks"
    parse_deck(path) -> the proposal deck (PPTX) as a list of slides

A "block" is one heading, paragraph, bullet or table, tagged with the RFP
section it belongs to. Keeping this structure (instead of one big string)
means later stages can say WHERE in the RFP a requirement came from.

No AI is used here: parsing is deterministic, fast, free and repeatable.
"""
import re
from pathlib import Path

import pdfplumber
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# ---------------------------------------------------------------------------
# Safety: the tool must never read the answer key (prevents data leakage)
# ---------------------------------------------------------------------------
BLOCKED_FOLDERS = {"ground_truth"}


def check_allowed(path):
    """Refuse files inside ground_truth/ and files that don't exist."""
    p = Path(path)
    if any(part in BLOCKED_FOLDERS for part in p.parts):
        raise PermissionError(f"Refusing to read '{p}': the tool must never read the answer key (ground_truth).")
    if not p.is_file():
        raise FileNotFoundError(f"File not found: '{p}'")
    return p


# ---------------------------------------------------------------------------
# Text cleaning
# ---------------------------------------------------------------------------
CID = re.compile(r"\(cid:\d+\)")          # PDF symbols pdfplumber can't decode, e.g. bullets -> "(cid:127)"
BULLET_START = re.compile(r"^\s*(\(cid:\d+\)|[•▪◦●\-–*])\s*")
NUMBERED_HEADING = re.compile(r"^\d+(\.\d+)*\.?\s+[A-Z]")   # e.g. "2. Scope of Work"


def clean(text, keep_lines=False):
    """Remove undecodable symbols, non-breaking spaces and repeated spaces.
    keep_lines=True keeps line breaks (used for table rows)."""
    text = CID.sub("", str(text)).replace("\u00a0", " ")
    if keep_lines:
        lines = (re.sub(r"[ \t]+", " ", ln).strip() for ln in text.splitlines())
        return "\n".join(ln for ln in lines if ln)
    return re.sub(r"\s+", " ", text).strip()


def _make_block(blocks, kind, text, section):
    text = clean(text, keep_lines=(kind == "table"))
    if text:
        blocks.append({"id": f"B{len(blocks) + 1:03d}", "type": kind, "section": section, "text": text})


def _finish(blocks, path, fmt):
    """Package blocks into one result, plus a readable full text for the LLM."""
    lines = []
    for b in blocks:
        prefix = {"heading": "## ", "bullet": "• "}.get(b["type"], "")
        lines.append(prefix + b["text"])
    full_text = "\n\n".join(lines)
    return {
        "source": Path(path).name,
        "format": fmt,
        "sections": [b["text"] for b in blocks if b["type"] == "heading"],
        "blocks": blocks,
        "full_text": full_text,
        "word_count": len(full_text.split()),
    }


# ---------------------------------------------------------------------------
# RFP: Word (.docx)
# ---------------------------------------------------------------------------
def _docx_items(doc):
    """Paragraphs AND tables, in the order they appear in the document."""
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def parse_docx(path):
    doc = Document(check_allowed(path))
    blocks, section = [], "Front matter"
    for item in _docx_items(doc):
        if isinstance(item, Table):
            rows = [" | ".join(clean(c.text) for c in row.cells) for row in item.rows]
            _make_block(blocks, "table", "\n".join(rows), section)
            continue
        style = item.style.name if item.style is not None else ""
        has_numbering = item._p.pPr is not None and item._p.pPr.numPr is not None
        if style.startswith("Heading"):
            section = clean(item.text) or section
            _make_block(blocks, "heading", item.text, section)
        elif style == "Title":
            _make_block(blocks, "paragraph", item.text, section)
        elif "List" in style or has_numbering:
            _make_block(blocks, "bullet", item.text, section)
        else:
            _make_block(blocks, "paragraph", item.text, section)
    return _finish(blocks, path, "docx")


# ---------------------------------------------------------------------------
# RFP: PDF
# ---------------------------------------------------------------------------
def parse_pdf(path):
    """
    PDFs store text as LINES, not paragraphs, so we rebuild them:
      - a numbered short line ("2. Scope of Work") starts a new section
      - a line starting with a bullet symbol starts a new bullet
      - a short line ending in . : ! ? closes the current paragraph
      - otherwise the line continues the current paragraph
    """
    with pdfplumber.open(check_allowed(path)) as pdf:
        lines = [ln for page in pdf.pages for ln in (page.extract_text() or "").splitlines() if ln.strip()]
    if not lines:
        raise ValueError(f"No text found in '{path}' (is it a scanned PDF?)")

    full_width = max(len(ln) for ln in lines)
    blocks, section = [], "Front matter"
    current, kind = [], "paragraph"

    def flush():
        nonlocal current
        if current:
            _make_block(blocks, kind, " ".join(current), section)
            current = []

    for ln in lines:
        text = ln.strip()
        if NUMBERED_HEADING.match(text) and len(text) < 80 and not text.endswith("."):
            flush()
            section = clean(text)
            _make_block(blocks, "heading", text, section)
            kind = "paragraph"
            continue
        if BULLET_START.match(text):
            flush()
            kind = "bullet"
            text = BULLET_START.sub("", text)
        current.append(text)
        if text.endswith((".", ":", "!", "?")) and len(ln) < 0.85 * full_width:
            flush()
            kind = "paragraph"
    flush()
    return _finish(blocks, path, "pdf")


def parse_rfp(path):
    """Read an RFP from .pdf or .docx (chosen by file extension)."""
    check_allowed(path)          # safety guard ALWAYS runs first, whatever the file type
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"Unsupported RFP format '{suffix}'. Use .pdf or .docx (save old .doc files as .docx).")


# ---------------------------------------------------------------------------
# Proposal deck: PowerPoint (.pptx)
# ---------------------------------------------------------------------------
def _walk_shapes(shapes):
    """All shapes in reading order (top to bottom, left to right), opening groups."""
    for shape in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)
        else:
            yield shape


def parse_deck(path):
    """
    One entry per slide: title, body lines, tables and speaker notes.
    'visible_text' = what the audience sees (title + body + tables).
    Speaker notes are kept separately: the client never sees them, so later
    stages can decide whether they count as "addressing" a requirement.
    """
    check_allowed(path)          # safety guard ALWAYS runs first
    if Path(path).suffix.lower() != ".pptx":
        raise ValueError(f"Unsupported deck format '{Path(path).suffix}'. Use .pptx (save old .ppt files as .pptx).")
    prs = Presentation(path)
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = clean(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        body, tables = [], []
        for shape in _walk_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_table:
                tables.append([[clean(c.text) for c in row.cells] for row in shape.table.rows])
            elif shape.has_text_frame:
                body.extend(t for t in (clean(p.text) for p in shape.text_frame.paragraphs) if t)
        notes = ""
        if slide.has_notes_slide:
            notes = clean(slide.notes_slide.notes_text_frame.text)

        parts = [f"Title: {title}"] + [f"- {b}" for b in body]
        for table in tables:
            parts += ["Table:"] + [" | ".join(row) for row in table]
        slides.append({
            "slide_number": number,
            "title": title,
            "body": body,
            "tables": tables,
            "notes": notes,
            "visible_text": "\n".join(parts),
        })
    return {"source": Path(path).name, "slide_count": len(slides), "slides": slides}
