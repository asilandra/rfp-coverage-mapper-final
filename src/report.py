"""
report.py - STAGE 5 of the tool: write the results to a colour-coded Excel file.

Sheets (answering the three questions in the brief):
  Summary     - totals (live Excel formulas), inputs, models used, how to read the file
  Coverage    - every requirement: status, slide(s), evidence quote, what's missing, reasoning
  Gaps        - only the requirements the deck does NOT fully address (brief item 3)
  Slide map   - for each slide: which requirements it addresses (brief item 2, from the slide side)
  Matrix      - requirements x slides grid (C = covers, P = partly covers)
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
STATUS_STYLE = {   # fill colour, font colour
    "Covered": ("C6EFCE", "006100"),
    "Partially covered": ("FFEB9C", "9C5700"),
    "Not covered": ("FFC7CE", "9C0006"),
    "Needs review": ("D9D9D9", "000000"),
}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def _status_cell(cell, status):
    fill, colour = STATUS_STYLE.get(status, ("FFFFFF", "000000"))
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = _font(color=colour, bold=True)


def _table(ws, headers, rows, widths, start_row=1):
    """Write a header row plus data rows, with wrapping, borders, filters and frozen header."""
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = _font()
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def _slides_text(numbers):
    return ", ".join(str(n) for n in numbers) if numbers else "-"


def write_report(verdicts, deck, rfp, path, judge_model="", retrieval_method="", unverified=None):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True     # Excel computes the Summary formulas when the file opens
    titles = {s["slide_number"]: s["title"] for s in deck["slides"]}
    n = len(verdicts)

    # ---------------- Coverage (one row per requirement) ----------------
    cov = wb.active
    cov.title = "Coverage"
    headers = ["ID", "Category", "Requirement (from RFP)", "Status", "Slide(s)", "Evidence (exact slide text)",
               "What's missing", "Reasoning", "Confidence", "Needs review", "RFP section", "RFP wording"]
    rows = [[v["id"], v["category"], v["requirement"], v["status"], _slides_text(v["slides"]), v["evidence"],
             v["missing"], v["reasoning"], v["confidence"], "Yes" if v["needs_review"] else "No",
             v["rfp_section"], v["rfp_quote"]] for v in verdicts]
    _table(cov, headers, rows, [6, 14, 45, 17, 10, 45, 38, 50, 11, 10, 26, 50])
    for r in range(2, n + 2):
        _status_cell(cov.cell(row=r, column=4), cov.cell(row=r, column=4).value)
        if cov.cell(row=r, column=10).value == "Yes":
            cov.cell(row=r, column=10).font = _font(bold=True, color="C00000")

    # ---------------- Gaps (not fully covered) ----------------
    gaps = wb.create_sheet("Gaps")
    gap_rows = [[v["id"], v["category"], v["requirement"], v["status"], v["missing"] or "-",
                 _slides_text(v["slides"]), v["rfp_section"], v["rfp_quote"]]
                for v in verdicts if v["status"] != "Covered"]
    _table(gaps, ["ID", "Category", "Requirement", "Status", "What's missing", "Slides that partly address it",
                  "RFP section", "RFP wording"], gap_rows, [6, 14, 45, 17, 50, 14, 26, 50])
    for r in range(2, len(gap_rows) + 2):
        _status_cell(gaps.cell(row=r, column=4), gaps.cell(row=r, column=4).value)
    if not gap_rows:
        gaps["A2"] = "No gaps: every requirement is fully covered."

    # ---------------- Slide map (from the slide side) ----------------
    smap = wb.create_sheet("Slide map")
    smap_rows = []
    for s in deck["slides"]:
        full = [v["id"] for v in verdicts if s["slide_number"] in v["slides"] and v["status"] == "Covered"]
        part = [v["id"] for v in verdicts if s["slide_number"] in v["slides"] and v["status"] == "Partially covered"]
        smap_rows.append([s["slide_number"], s["title"], ", ".join(full) or "-", ", ".join(part) or "-",
                          len(full) + len(part)])
    _table(smap, ["Slide", "Slide title", "Fully addresses", "Partly addresses", "Number of requirements"],
           smap_rows, [7, 40, 35, 35, 14])

    # ---------------- Matrix (requirements x slides) ----------------
    mx = wb.create_sheet("Matrix")
    slide_numbers = [s["slide_number"] for s in deck["slides"]]
    mx_headers = ["ID", "Requirement", "Status"] + [f"S{k}" for k in slide_numbers]
    mx_rows = []
    for v in verdicts:
        mark = "C" if v["status"] == "Covered" else "P"
        mx_rows.append([v["id"], v["requirement"], v["status"]] +
                       [mark if k in v["slides"] else "" for k in slide_numbers])
    _table(mx, mx_headers, mx_rows, [6, 50, 17] + [5] * len(slide_numbers))
    for r in range(2, n + 2):
        _status_cell(mx.cell(row=r, column=3), mx.cell(row=r, column=3).value)
        for c in range(4, 4 + len(slide_numbers)):
            cell = mx.cell(row=r, column=c)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.value == "C":
                _status_cell(cell, "Covered")
            elif cell.value == "P":
                _status_cell(cell, "Partially covered")
    for c, k in enumerate(slide_numbers, start=4):   # slide titles as header comments
        from openpyxl.comments import Comment
        mx.cell(row=1, column=c).comment = Comment(f"Slide {k}: {titles[k]}", "Coverage Mapper")

    # ---------------- Summary (first sheet) ----------------
    sm = wb.create_sheet("Summary", 0)
    sm.column_dimensions["A"].width = 34
    sm.column_dimensions["B"].width = 40
    sm["A1"] = "RFP vs Proposal Coverage Report"
    sm["A1"].font = _font(size=14, bold=True, color="1F3864")
    info = [("RFP", rfp["source"]), ("Proposal deck", f"{deck['source']} ({deck['slide_count']} slides)"),
            ("Report date", date.today().isoformat()), ("Judge model (all requirements)", judge_model),
            ("Slide ranking", retrieval_method)]
    for i, (k, v) in enumerate(info, start=3):
        sm.cell(row=i, column=1, value=k).font = _font(bold=True)
        sm.cell(row=i, column=2, value=v).font = _font()

    sm["A9"] = "Totals (live formulas from the Coverage sheet)"
    sm["A9"].font = _font(bold=True, size=11, color="1F3864")
    last = n + 1
    totals = [
        ("Requirements found in the RFP", f"=COUNTA(Coverage!A2:A{last})", None),
        ("Covered", f'=COUNTIF(Coverage!D2:D{last},"Covered")', "Covered"),
        ("Partially covered", f'=COUNTIF(Coverage!D2:D{last},"Partially covered")', "Partially covered"),
        ("Not covered", f'=COUNTIF(Coverage!D2:D{last},"Not covered")', "Not covered"),
        ("Gaps to fix (partial + not covered)", "=B12+B13", None),
        ("Fully covered (%)", "=IFERROR(B11/B10,0)", None),
        ("Verdicts flagged for human review", f'=COUNTIF(Coverage!J2:J{last},"Yes")', None),
    ]
    for i, (label, formula, style) in enumerate(totals, start=10):
        a, b = sm.cell(row=i, column=1, value=label), sm.cell(row=i, column=2, value=formula)
        a.font, b.font = _font(), _font(bold=True)
        b.alignment = Alignment(horizontal="left")
        if style:
            _status_cell(a, style)
    sm["B15"].number_format = "0%"

    sm["A18"] = "By category"
    sm["A18"].font = _font(bold=True, size=11, color="1F3864")
    for c, h in enumerate(["Category", "Covered", "Partially covered", "Not covered"], start=1):
        cell = sm.cell(row=19, column=c, value=h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for col in ("C", "D"):
        sm.column_dimensions[col].width = 18
    for i, cat in enumerate(["Scope", "Deliverable", "Timeline", "Proposal requirement"], start=20):
        sm.cell(row=i, column=1, value=cat).font = _font()
        for c, status in enumerate(["Covered", "Partially covered", "Not covered"], start=2):
            cell = sm.cell(row=i, column=c,
                           value=f'=COUNTIFS(Coverage!B2:B{last},A{i},Coverage!D2:D{last},"{status}")')
            cell.font = _font()
            cell.alignment = Alignment(horizontal="center")

    sm["A26"] = "How to read this report"
    sm["A26"].font = _font(bold=True, size=11, color="1F3864")
    notes = [
        "Coverage: one row per RFP requirement, with the slide(s) that address it and an exact quote from the slide as evidence.",
        "Gaps: requirements the deck does not fully address, and what is missing. Start here before submission.",
        "Slide map / Matrix: the same results seen from the slides' side (C = covers, P = partly covers).",
        "Covered = the slides commit to every part of the requirement. Partially covered = at least one part is missing. Not covered = no slide commits to it.",
        "Every requirement quotes the RFP word for word, and every evidence quote was checked by code to exist on the cited slide. 'Needs review' marks verdicts that failed a check.",
        "Speaker notes are not counted: the client never sees them.",
    ]
    for i, text in enumerate(notes, start=27):
        cell = sm.cell(row=i, column=1, value=f"• {text}")
        cell.font = _font()
        sm.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sm.row_dimensions[i].height = 24

    if unverified:
        rv = wb.create_sheet("Set aside")
        _table(rv, ["Category", "Requirement proposed by the AI", "Quote it gave", "Why set aside"],
               [[u["category"], u["requirement"], u["source_quote"], "Quote not found in the RFP"] for u in unverified],
               [14, 45, 50, 25])

    for ws in wb.worksheets:                 # print-friendly: landscape, fit all columns on the page width
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(path)
    return path
