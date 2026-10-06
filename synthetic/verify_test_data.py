"""
verify_test_data.py - checks the generated test set is complete and correct.

Run from the project root:
    python -m synthetic.verify_test_data

It opens the REAL files (not Gemini's raw JSON), so it checks what the tool
will actually see. Every check prints a tick or a cross.
"""
import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")   # Windows: print "?" instead of crashing on special characters

import json
from pathlib import Path

import pdfplumber
from docx import Document
from pptx import Presentation

from synthetic.scenario import REQUIREMENTS, SLIDES

INPUTS = Path("test_data/inputs")
TRUTH = Path("test_data/ground_truth")
results = []


def check(ok, label):
    results.append(ok)
    print(("  ✅ " if ok else "  ❌ ") + label)
    return ok


def slide_text(slide):
    """All text on a slide: text boxes, tables and speaker notes."""
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            for row in shape.table.rows:
                parts.extend(cell.text for cell in row.cells)
    if slide.has_notes_slide:
        parts.append(slide.notes_slide.notes_text_frame.text)
    return " ".join(parts)


print("1) All files exist")
files = [INPUTS / "rfp_qamar.docx", INPUTS / "rfp_qamar.pdf", INPUTS / "proposal_qamar.pptx",
         TRUTH / "answer_key.json", TRUTH / "answer_key.csv",
         TRUTH / "rfp_content.json", TRUTH / "deck_content.json"]
for f in files:
    check(f.exists() and f.stat().st_size > 0, str(f))
if not all(results):
    raise SystemExit("\nFiles are missing - run the 'Generate test data' cell first.")

print("\n2) RFP is readable and complete")
pdf_text = "\n".join(page.extract_text() or "" for page in pdfplumber.open(INPUTS / "rfp_qamar.pdf").pages)
docx_words = sum(len(p.text.split()) for p in Document(INPUTS / "rfp_qamar.docx").paragraphs)
check(len(pdf_text.split()) > 500, f"PDF text readable ({len(pdf_text.split())} words)")
check(docx_words > 500, f"Word text readable ({docx_words} words)")
locations = json.loads((TRUTH / "rfp_content.json").read_text(encoding="utf-8")).get("requirement_locations", {})
missing = [r["id"] for r in REQUIREMENTS if r["id"] not in locations]
check(not missing, "Gemini placed all 18 requirements in the RFP" + (f" (missing: {missing})" if missing else ""))

print("\n3) Deck structure")
prs = Presentation(INPUTS / "proposal_qamar.pptx")
slides = list(prs.slides)
check(len(slides) == len(SLIDES), f"{len(slides)} slides (expected {len(SLIDES)})")
check(any(sh.has_table for sh in slides[10].shapes) if len(slides) > 10 else False, "Slide 11 contains a table")
with_notes = sum(1 for s in slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
print(f"  ℹ️ {with_notes}/{len(slides)} slides have speaker notes (information only)")

print("\n4) Planted gaps are intact (no forbidden words in the real slides)")
for plan, slide in zip(SLIDES, slides):
    if plan.get("avoid"):
        text = slide_text(slide).lower()
        used = [w for w in plan["avoid"] if w.lower() in text]
        check(not used, f"Slide {plan['n']} ({plan['title']})" + (f" uses: {used}" if used else ""))

print("\n5) Answer key")
key = json.loads((TRUTH / "answer_key.json").read_text(encoding="utf-8"))
counts = {s: sum(r["expected_status"] == s for r in key) for s in ("Covered", "Partially covered", "Not covered")}
check(len(key) == 18, f"{len(key)} requirements in the answer key")
check(counts == {"Covered": 10, "Partially covered": 4, "Not covered": 4}, f"Status split: {counts}")

failed = results.count(False)
print("\n" + ("🎉 ALL CHECKS PASSED - the test data is complete." if failed == 0
              else f"⚠️ {failed} check(s) failed - paste this output to Claude."))
